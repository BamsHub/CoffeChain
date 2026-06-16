import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def main():
    json_path = 'perf_results.json'
    excel_path = 'Blockchain_Performance_Report.xlsx'

    if not os.path.exists(json_path):
        print(f"Error: {json_path} not found. Please run speed test first.")
        return

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # Create workbook and sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Performance Report"
    ws.views.sheetView[0].showGridLines = True

    # Styling Definitions
    font_title = Font(name='Segoe UI', size=16, bold=True, color='1B4332') # Forest Green
    font_section = Font(name='Segoe UI', size=12, bold=True, color='2D6A4F')
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Segoe UI', size=11, bold=True, color='000000')
    font_regular = Font(name='Segoe UI', size=11, color='333333')
    font_link = Font(name='Segoe UI', size=10, color='0A9396', underline='single')
    
    fill_header = PatternFill(start_color='2D6A4F', end_color='2D6A4F', fill_type='solid') # Forest Green Header
    fill_card = PatternFill(start_color='D8F3DC', end_color='D8F3DC', fill_type='solid') # Pale Green
    fill_zebra = PatternFill(start_color='F4F9F4', end_color='F4F9F4', fill_type='solid') # Very light zebra row
    fill_success = PatternFill(start_color='A3E635', end_color='A3E635', fill_type='solid') # Soft Lime Green
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    # 1. Title Block
    ws['B2'] = "☕ CoffeeChain Blockchain Performance Report"
    ws['B2'].font = font_title
    ws['B2'].alignment = align_left
    
    ws['B3'] = "Solana Testnet Transaction Latency and Gas Costs Analysis"
    ws['B3'].font = Font(name='Segoe UI', size=11, italic=True, color='52B788')
    ws['B3'].alignment = align_left
    
    # 2. Key Metrics Summary Cards
    metrics = [
        ("Total Transactions", f"{data['summary']['total']} TXs"),
        ("Success Rate", f"{int(data['summary']['success'] / data['summary']['total'] * 100)}%"),
        ("Avg Latency", f"{data['summary']['avg_seconds']} seconds"),
        ("Total Gas Fee", f"{data['summary']['cost_sol']} SOL")
    ]
    
    # Write summary cards side by side
    col_offsets = ['B', 'D', 'F', 'H']
    for idx, (label, val) in enumerate(metrics):
        col = col_offsets[idx]
        next_col = chr(ord(col) + 1)
        
        # Merge cells for card
        ws.merge_cells(f"{col}5:{next_col}6")
        cell = ws[f"{col}5"]
        cell.value = f"{label}\n{val}"
        cell.font = font_bold
        cell.fill = fill_card
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Border for the merged card
        for r in range(5, 7):
            for c in [col, next_col]:
                ws[f"{c}{r}"].border = border_thin

    # 3. Transaction Details Table Section
    ws['B8'] = "📋 Detailed Transaction Log"
    ws['B8'].font = font_section
    ws['B8'].alignment = align_left
    
    headers = ["Tx #", "Status", "Latency (s)", "Transaction Signature", "Explorer Link"]
    start_row = 9
    start_col = 2 # Column B
    
    # Write Headers
    for c_idx, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + c_idx)
        cell.value = header
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    # Write Data
    for r_idx, tx in enumerate(data['transactions']):
        row = start_row + 1 + r_idx
        is_zebra = r_idx % 2 == 1
        row_fill = fill_zebra if is_zebra else PatternFill(fill_type=None)
        
        # Tx No
        c_tx = ws.cell(row=row, column=start_col)
        c_tx.value = tx['tx']
        c_tx.font = font_regular
        c_tx.fill = row_fill
        c_tx.alignment = align_center
        c_tx.border = border_thin
        
        # Status
        c_status = ws.cell(row=row, column=start_col + 1)
        c_status.value = tx['status']
        c_status.font = font_bold
        c_status.fill = fill_success if tx['status'] == 'OK' else row_fill
        c_status.alignment = align_center
        c_status.border = border_thin
        
        # Latency
        c_latency = ws.cell(row=row, column=start_col + 2)
        c_latency.value = tx['seconds']
        c_latency.font = font_regular
        c_latency.fill = row_fill
        c_latency.alignment = align_right
        c_latency.border = border_thin
        
        # Signature
        c_sig = ws.cell(row=row, column=start_col + 3)
        c_sig.value = tx['signature']
        c_sig.font = font_regular
        c_sig.fill = row_fill
        c_sig.alignment = align_left
        c_sig.border = border_thin
        
        # Explorer Link (Hyperlink)
        c_link = ws.cell(row=row, column=start_col + 4)
        c_link.value = "Solana Explorer ↗"
        c_link.hyperlink = f"https://explorer.solana.com/tx/{tx['signature']}?cluster=testnet"
        c_link.font = font_link
        c_link.fill = row_fill
        c_link.alignment = align_center
        c_link.border = border_thin

    # Auto-adjust column widths
    for col in ws.columns:
        # Check if it's outside our used range
        col_letter = get_column_letter(col[0].column)
        if col_letter < 'B' or col_letter > 'F':
            continue
        max_len = 0
        for cell in col:
            # Avoid title cell spanning multiple column sizes
            if cell.row < 8:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Explicit override for Signature column width
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 20

    # Save to file
    wb.save(excel_path)
    print(f"Beautiful Excel performance report generated successfully at: {excel_path}")

if __name__ == '__main__':
    main()
